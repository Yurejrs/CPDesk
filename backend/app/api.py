import csv
import io
from datetime import date, datetime, timedelta
from decimal import Decimal

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from .database import get_db
from .models import (FuelRecord, Item, Maintenance, PurchaseSchedule, Sector,
                     StockMovement, User, Vehicle)
from .schemas import (FuelInput, ItemInput, LoginInput, MaintenanceInput, MovementInput,
                      ReasonInput, ScheduleCompleteInput, ScheduleInput, SectorInput,
                      UserCreate, VehicleInput)
from .security import admin_user, create_token, current_user, hash_password, verify_password
from .services import (audit, dashboard_data, fuel_efficiency, item_balance, item_balances,
                       report_data, stock_alerts, stock_consumption, validate_odometer)

router = APIRouter(prefix="/api/v1")


@router.get("/auth/setup-status")
def setup_status(db: Session = Depends(get_db)):
    return {"needs_setup": db.scalar(select(User.id).limit(1)) is None}


@router.post("/auth/setup", status_code=201)
def initial_setup(payload: UserCreate, response: Response, db: Session = Depends(get_db)):
    if db.scalar(select(User.id).limit(1)):
        raise HTTPException(409, "Configuração inicial já concluída")
    if payload.role != "admin":
        raise HTTPException(422, "O primeiro usuário precisa ser administrador")
    user = User(username=payload.username, full_name=payload.full_name,
                password_hash=hash_password(payload.password), role="admin")
    db.add(user)
    db.commit()
    db.refresh(user)
    token = create_token(user)
    response.set_cookie("cpdesk_session", token, httponly=True, samesite="lax", max_age=43200)
    return {"user": {"id": user.id, "username": user.username, "full_name": user.full_name, "role": user.role}, "access_token": token}


@router.post("/auth/login")
def login(payload: LoginInput, response: Response, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.username == payload.username))
    if not user or not user.is_active or not verify_password(payload.password, user.password_hash):
        raise HTTPException(401, "Usuário ou senha inválidos")
    token = create_token(user)
    response.set_cookie("cpdesk_session", token, httponly=True, samesite="lax", max_age=43200)
    return {"access_token": token, "user": {"id": user.id, "username": user.username, "full_name": user.full_name, "role": user.role}}


@router.post("/auth/logout")
def logout(response: Response):
    response.delete_cookie("cpdesk_session")
    return {"ok": True}


@router.get("/auth/me")
def me(user: User = Depends(current_user)):
    return {"id": user.id, "username": user.username, "full_name": user.full_name, "role": user.role}


@router.post("/users", status_code=201)
def create_user(payload: UserCreate, db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    if db.scalar(select(User.id).where(User.username == payload.username)):
        raise HTTPException(409, "Nome de usuário já cadastrado")
    user = User(username=payload.username, full_name=payload.full_name,
                password_hash=hash_password(payload.password), role=payload.role)
    db.add(user)
    db.flush()
    audit(db, admin.id, "user.created", "user", user.id)
    db.commit()
    return {"id": user.id, "username": user.username, "full_name": user.full_name, "role": user.role, "is_active": user.is_active}


@router.get("/users")
def list_users(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    return [{"id": u.id, "username": u.username, "full_name": u.full_name, "role": u.role, "is_active": u.is_active}
            for u in db.scalars(select(User).order_by(User.full_name)).all()]


@router.get("/sectors")
def list_sectors(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    return db.scalars(select(Sector).where(Sector.is_active.is_(True)).order_by(Sector.name)).all()


@router.post("/sectors", status_code=201)
def create_sector(payload: SectorInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    if db.scalar(select(Sector.id).where(Sector.name == payload.name.strip())):
        raise HTTPException(409, "Setor já cadastrado")
    sector = Sector(name=payload.name.strip())
    db.add(sector)
    db.flush()
    audit(db, user.id, "sector.created", "sector", sector.id)
    db.commit()
    return sector


@router.get("/items")
def list_items(q: str | None = None, db: Session = Depends(get_db), _user: User = Depends(current_user)):
    query = select(Item).where(Item.is_active.is_(True)).order_by(Item.name)
    if q:
        query = query.where(Item.name.ilike(f"%{q}%") | Item.sku.ilike(f"%{q}%"))
    balances = item_balances(db)
    return [{"id": item.id, "sku": item.sku, "name": item.name, "category": item.category,
             "unit": item.unit, "minimum_stock": item.minimum_stock, "current_unit_cost": item.current_unit_cost,
             "last_supplier": item.last_supplier, "is_active": item.is_active,
             "balance": balances.get(item.id, Decimal("0")),
             "low_stock": item.minimum_stock > 0 and balances.get(item.id, Decimal("0")) <= item.minimum_stock}
            for item in db.scalars(query).all()]


@router.post("/items", status_code=201)
def create_item(payload: ItemInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    if payload.sku and db.scalar(select(Item.id).where(Item.sku == payload.sku)):
        raise HTTPException(409, "Código do item já cadastrado")
    item = Item(**(payload.model_dump() | {"name": payload.name.strip()}))
    db.add(item)
    db.flush()
    audit(db, user.id, "item.created", "item", item.id)
    db.commit()
    return {"id": item.id, "name": item.name, "balance": Decimal("0"), "low_stock": item.minimum_stock > 0}


@router.post("/items/{item_id}/opening-balance", status_code=201)
def set_opening_balance(item_id: int, payload: MovementInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    item = db.get(Item, item_id)
    if not item:
        raise HTTPException(404, "Item não encontrado")
    if payload.movement_type != "entry":
        raise HTTPException(422, "Saldo inicial deve ser uma entrada")
    if item_balance(db, item_id) != 0:
        raise HTTPException(409, "Saldo inicial só pode ser informado antes de outras movimentações")
    payload.reason = f"Saldo inicial: {payload.reason}"
    movement = _record_movement(payload, item, db, user)
    db.commit()
    db.refresh(movement)
    return movement


@router.post("/items/import", status_code=201)
async def import_items(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(current_user)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(415, "Envie uma planilha Excel (.xlsx)")
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
    except Exception as exc:
        raise HTTPException(422, "Não foi possível ler a planilha Excel") from exc
    if not values or not values[0]:
        raise HTTPException(422, "A planilha está vazia")
    def normalize_header(value):
        import unicodedata
        text = "".join(ch for ch in unicodedata.normalize("NFKD", str(value).strip().lower())
                      if not unicodedata.combining(ch))
        return " ".join(text.split())

    headers = {normalize_header(value): index for index, value in enumerate(values[0]) if value is not None}
    name_index = headers.get("descricao", headers.get("nome", headers.get("item")))
    if name_index is None:
        raise HTTPException(422, "A planilha precisa de uma coluna Descrição (ou Nome)")

    def cell(row, *names):
        index = next((headers[name] for name in names if name in headers), None)
        return row[index] if index is not None and index < len(row) else None

    imported = []
    try:
        for line_number, row in enumerate(values[1:], start=2):
            if not row or not any(value is not None for value in row):
                continue
            name = str(row[name_index]).strip() if row[name_index] is not None else ""
            if not name:
                raise ValueError(f"Descrição vazia na linha {line_number}")
            sku = cell(row, "codigo", "sku")
            sku = str(sku).strip() if sku is not None else None
            if sku and db.scalar(select(Item.id).where(Item.sku == sku)):
                raise ValueError(f"Código {sku} já cadastrado (linha {line_number})")
            if sku and any(existing.sku == sku for existing in imported):
                raise ValueError(f"Código {sku} duplicado na planilha (linha {line_number})")
            minimum = Decimal(str(cell(row, "estoque minimo", "minimo", "minimum_stock") or 0))
            unit_cost_raw = cell(row, "valor unitario", "custo unitario")
            unit_cost = Decimal(str(unit_cost_raw)) if unit_cost_raw not in (None, "") else None
            item = Item(sku=sku, name=name, category=str(cell(row, "categoria") or "") or None,
                unit=str(cell(row, "unidade", "unit") or "un"), minimum_stock=minimum,
                current_unit_cost=unit_cost, last_supplier=str(cell(row, "fornecedor") or "") or None)
            db.add(item)
            db.flush()
            opening = cell(row, "saldo inicial", "saldo", "quantidade")
            if opening not in (None, "", 0):
                quantity = Decimal(str(opening))
                if quantity < 0:
                    raise ValueError(f"Saldo inicial negativo na linha {line_number}")
                if quantity > 0:
                    movement = StockMovement(item_id=item.id, movement_type="entry", quantity=quantity,
                        signed_quantity=quantity, reason="Importação de planilha: saldo inicial",
                        unit_cost=unit_cost, invoice_total=unit_cost * quantity if unit_cost is not None else None,
                        supplier_name=item.last_supplier, responsible_id=user.id)
                    db.add(movement)
            imported.append(item)
        audit(db, user.id, "items.imported", "item", imported[-1].id if imported else 0,
            f"items={len(imported)};filename={file.filename}")
        db.commit()
    except Exception as exc:
        db.rollback()
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(422, f"Importação cancelada: {exc}") from exc
    return {"imported": len(imported)}


@router.patch("/items/{item_id}")
def update_item(item_id: int, payload: ItemInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    item = db.get(Item, item_id)
    if not item:
        raise HTTPException(404, "Item não encontrado")
    if payload.sku and db.scalar(select(Item.id).where(Item.sku == payload.sku, Item.id != item_id)):
        raise HTTPException(409, "Código do item já cadastrado")
    for key, value in payload.model_dump().items():
        setattr(item, key, value)
    audit(db, user.id, "item.updated", "item", item.id)
    db.commit()
    return {"id": item.id, "name": item.name, "balance": item_balance(db, item.id), "low_stock": item.minimum_stock > 0 and item_balance(db, item.id) <= item.minimum_stock}


@router.get("/stock/alerts")
def get_stock_alerts(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    return stock_alerts(db)


@router.get("/stock/movements")
def list_movements(item_id: int | None = None, db: Session = Depends(get_db), _user: User = Depends(current_user)):
    query = select(StockMovement).order_by(StockMovement.created_at.desc(), StockMovement.id.desc())
    if item_id:
        query = query.where(StockMovement.item_id == item_id)
    movements = db.scalars(query).all()
    return [{"id": row.id, "item_id": row.item_id, "movement_type": row.movement_type,
        "quantity": row.quantity, "signed_quantity": row.signed_quantity, "reason": row.reason,
        "sector_id": row.sector_id, "sector_name": row.sector.name if row.sector else None,
        "responsible_id": row.responsible_id, "responsible_name": row.responsible.full_name if row.responsible else None,
        "supplier_name": row.supplier_name, "invoice_number": row.invoice_number,
        "unit_cost": row.unit_cost, "invoice_total": row.invoice_total,
        "reverses_id": row.reverses_id, "created_at": row.created_at} for row in movements]


def _record_movement(payload: MovementInput, item: Item, db: Session, user: User):
    if payload.movement_type not in {"entry", "exit", "adjustment"}:
        raise HTTPException(422, "Tipo de movimentação inválido")
    if payload.movement_type == "exit" and item_balance(db, item.id) < payload.quantity:
        raise HTTPException(409, "Saldo insuficiente")
    if payload.movement_type == "entry":
        signed = payload.quantity
    elif payload.movement_type == "exit":
        signed = -payload.quantity
    else:
        if payload.signed_quantity is None or payload.signed_quantity == 0:
            raise HTTPException(422, "Ajuste exige uma diferença assinada diferente de zero")
        signed = payload.signed_quantity
    if item_balance(db, item.id) + signed < 0:
        raise HTTPException(409, "A movimentação deixaria o saldo negativo")
    if payload.movement_type == "exit" and (not payload.sector_id or not payload.responsible_id):
        raise HTTPException(422, "Consumo exige setor e funcionário responsável")
    if payload.movement_type == "exit" and not db.get(Sector, payload.sector_id):
        raise HTTPException(422, "Setor não encontrado")
    if payload.sector_id and not db.get(Sector, payload.sector_id):
        raise HTTPException(422, "Setor não encontrado")
    responsible_id = payload.responsible_id or user.id
    if not db.get(User, responsible_id):
        raise HTTPException(422, "Responsável não encontrado")
    unit_cost = payload.unit_cost if payload.unit_cost is not None else item.current_unit_cost
    invoice_total = payload.invoice_total
    if payload.movement_type == "entry":
        if not payload.supplier_name or invoice_total is None:
            raise HTTPException(422, "Entrada exige fornecedor e valor total da nota fiscal")
        if unit_cost is None:
            unit_cost = invoice_total / payload.quantity
        if unit_cost is not None:
            item.current_unit_cost = unit_cost
        if payload.supplier_name:
            item.last_supplier = payload.supplier_name.strip()
    movement = StockMovement(item_id=item.id, movement_type=payload.movement_type, quantity=payload.quantity,
        signed_quantity=signed, reason=payload.reason, sector_id=payload.sector_id,
        responsible_id=responsible_id, unit_cost=unit_cost, invoice_total=invoice_total,
        supplier_name=payload.supplier_name.strip() if payload.supplier_name else None,
        invoice_number=payload.invoice_number)
    db.add(movement)
    db.flush()
    audit(db, user.id, "stock.movement.created", "stock_movement", movement.id, payload.reason)
    return movement


@router.post("/items/{item_id}/movements", status_code=201)
def create_movement(item_id: int, payload: MovementInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    item = db.get(Item, item_id)
    if not item:
        raise HTTPException(404, "Item não encontrado")
    movement = _record_movement(payload, item, db, user)
    db.commit()
    db.refresh(movement)
    return movement


@router.post("/stock/movements/{movement_id}/reverse", status_code=201)
def reverse_movement(movement_id: int, payload: ReasonInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    movement = db.get(StockMovement, movement_id)
    if not movement:
        raise HTTPException(404, "Movimentação não encontrada")
    if movement.reverses_id or db.scalar(select(StockMovement.id).where(StockMovement.reverses_id == movement.id)):
        raise HTTPException(409, "Movimentação já estornada")
    if movement.signed_quantity > 0 and item_balance(db, movement.item_id) < movement.signed_quantity:
        raise HTTPException(409, "Não é possível estornar: o saldo atual já foi consumido")
    reversal = StockMovement(item_id=movement.item_id, movement_type="adjustment", quantity=movement.quantity,
        signed_quantity=-movement.signed_quantity, reason=f"Estorno de #{movement.id}: {payload.reason}",
        responsible_id=user.id, sector_id=movement.sector_id, reverses_id=movement.id, unit_cost=movement.unit_cost,
        invoice_total=movement.invoice_total, supplier_name=movement.supplier_name,
        invoice_number=movement.invoice_number)
    db.add(reversal)
    db.flush()
    audit(db, user.id, "stock.movement.reversed", "stock_movement", movement.id, payload.reason)
    db.commit()
    return reversal


@router.get("/purchases")
def list_purchases(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    today = date.today()
    rows = db.scalars(select(PurchaseSchedule).where(PurchaseSchedule.is_active.is_(True))
        .order_by(PurchaseSchedule.next_due_date, PurchaseSchedule.supplier_name)).all()
    return [{"id": row.id, "supplier_name": row.supplier_name, "recurrence_days": row.recurrence_days,
        "next_due_date": row.next_due_date, "last_completed_date": row.last_completed_date,
        "days_until_due": (row.next_due_date - today).days, "notes": row.notes}
        for row in rows]


@router.post("/purchases", status_code=201)
def create_purchase_schedule(payload: ScheduleInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    row = PurchaseSchedule(supplier_name=payload.supplier_name.strip(), recurrence_days=payload.recurrence_days,
        next_due_date=payload.next_due_date, notes=payload.notes, created_by=user.id)
    db.add(row)
    db.flush()
    audit(db, user.id, "purchase_schedule.created", "purchase_schedule", row.id)
    db.commit()
    return {"id": row.id, "supplier_name": row.supplier_name, "next_due_date": row.next_due_date}


@router.patch("/purchases/{schedule_id}")
def update_purchase_schedule(schedule_id: int, payload: ScheduleInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    row = db.get(PurchaseSchedule, schedule_id)
    if not row:
        raise HTTPException(404, "Rotina de compra não encontrada")
    for key, value in payload.model_dump().items():
        setattr(row, key, value)
    audit(db, user.id, "purchase_schedule.updated", "purchase_schedule", row.id)
    db.commit()
    return {"id": row.id, "supplier_name": row.supplier_name, "next_due_date": row.next_due_date}


@router.post("/purchases/{schedule_id}/complete")
def complete_purchase_schedule(schedule_id: int, payload: ScheduleCompleteInput,
                               db: Session = Depends(get_db), user: User = Depends(current_user)):
    row = db.get(PurchaseSchedule, schedule_id)
    if not row or not row.is_active:
        raise HTTPException(404, "Rotina de compra não encontrada ou arquivada")
    if payload.completed_on < row.next_due_date:
        raise HTTPException(422, "A rotina não pode ser concluída antes da data prevista")
    row.last_completed_date = payload.completed_on
    next_due = payload.completed_on + timedelta(days=row.recurrence_days)
    if next_due.weekday() == 5:
        next_due += timedelta(days=2)
    elif next_due.weekday() == 6:
        next_due += timedelta(days=1)
    row.next_due_date = next_due
    audit(db, user.id, "purchase_schedule.completed", "purchase_schedule", row.id,
          f"completed_on={payload.completed_on.isoformat()};next_due={next_due.isoformat()}")
    db.commit()
    return {"id": row.id, "last_completed_date": row.last_completed_date, "next_due_date": row.next_due_date}


@router.post("/purchases/{schedule_id}/archive")
def archive_purchase_schedule(schedule_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    row = db.get(PurchaseSchedule, schedule_id)
    if not row:
        raise HTTPException(404, "Rotina de compra não encontrada")
    row.is_active = False
    audit(db, user.id, "purchase_schedule.archived", "purchase_schedule", row.id)
    db.commit()
    return {"id": row.id, "is_active": False}


@router.get("/vehicles")
def list_vehicles(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    return db.scalars(select(Vehicle).order_by(Vehicle.name)).all()


def normalize_plate(value: str) -> str:
    import re
    return re.sub(r"[^A-Z0-9]", "", value.upper())


@router.post("/vehicles", status_code=201)
def create_vehicle(payload: VehicleInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    plate = normalize_plate(payload.plate)
    if len(plate) != 7:
        raise HTTPException(422, "Placa deve conter sete caracteres")
    if db.scalar(select(Vehicle.id).where(Vehicle.plate == plate)):
        raise HTTPException(409, "Placa já cadastrada")
    vehicle = Vehicle(**payload.model_dump(exclude={"plate"}), plate=plate)
    db.add(vehicle)
    db.flush()
    audit(db, user.id, "vehicle.created", "vehicle", vehicle.id)
    db.commit()
    return vehicle


@router.post("/vehicles/{vehicle_id}/fuel", status_code=201)
def add_fuel(vehicle_id: int, payload: FuelInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    vehicle = db.get(Vehicle, vehicle_id)
    if not vehicle or not vehicle.is_active:
        raise HTTPException(404, "Veículo não encontrado ou inativo")
    violations = validate_odometer(db, vehicle, payload.occurred_on, payload.odometer)
    warning = "; ".join(violations) if violations else None
    if violations:
        audit(db, user.id, "fuel.odometer.warning", "vehicle", vehicle.id, warning)
    record = FuelRecord(vehicle_id=vehicle.id, occurred_on=payload.occurred_on, odometer=payload.odometer,
        liters=payload.liters, amount=payload.amount, full_tank=payload.full_tank, fuel_type=payload.fuel_type,
        station=payload.station, driver=payload.driver, odometer_warning=warning)
    db.add(record)
    db.flush()
    audit(db, user.id, "fuel.created", "fuel_record", record.id)
    db.commit()
    return record


@router.post("/vehicles/{vehicle_id}/fuel/{record_id}/cancel")
def cancel_fuel(vehicle_id: int, record_id: int, payload: ReasonInput,
                db: Session = Depends(get_db), user: User = Depends(admin_user)):
    record = db.get(FuelRecord, record_id)
    if not record or record.vehicle_id != vehicle_id:
        raise HTTPException(404, "Abastecimento não encontrado")
    if record.cancelled_at:
        raise HTTPException(409, "Abastecimento já cancelado")
    record.cancelled_at = datetime.utcnow()
    record.cancellation_reason = payload.reason
    audit(db, user.id, "fuel.cancelled", "fuel_record", record.id, payload.reason)
    db.commit()
    return {"id": record.id, "cancelled": True}


@router.get("/vehicles/{vehicle_id}/fuel")
def vehicle_fuel_history(vehicle_id: int, db: Session = Depends(get_db), _user: User = Depends(current_user)):
    records = db.scalars(select(FuelRecord).where(FuelRecord.vehicle_id == vehicle_id,
        FuelRecord.cancelled_at.is_(None)).order_by(FuelRecord.occurred_on, FuelRecord.odometer, FuelRecord.id)).all()
    if not records and not db.get(Vehicle, vehicle_id):
        raise HTTPException(404, "Veículo não encontrado")
    return fuel_efficiency(records)


@router.post("/vehicles/{vehicle_id}/maintenance", status_code=201)
def add_maintenance(vehicle_id: int, payload: MaintenanceInput, db: Session = Depends(get_db), user: User = Depends(current_user)):
    vehicle = db.get(Vehicle, vehicle_id)
    if not vehicle:
        raise HTTPException(404, "Veículo não encontrado")
    row = Maintenance(vehicle_id=vehicle_id, **payload.model_dump())
    db.add(row)
    db.flush()
    audit(db, user.id, "maintenance.created", "maintenance", row.id)
    db.commit()
    return row


@router.get("/vehicles/{vehicle_id}/maintenance")
def maintenance_history(vehicle_id: int, db: Session = Depends(get_db), _user: User = Depends(current_user)):
    if not db.get(Vehicle, vehicle_id):
        raise HTTPException(404, "Veículo não encontrado")
    return db.scalars(select(Maintenance).where(Maintenance.vehicle_id == vehicle_id).order_by(Maintenance.occurred_on.desc())).all()


@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    return dashboard_data(db)


@router.get("/reports/fleet")
def fleet_report(start: date | None = None, end: date | None = None,
                 db: Session = Depends(get_db), _user: User = Depends(current_user)):
    if start and end and start > end:
        raise HTTPException(422, "Data inicial posterior à data final")
    return report_data(db, start, end)


@router.get("/reports/fleet.csv")
def fleet_report_csv(start: date | None = None, end: date | None = None,
                     db: Session = Depends(get_db), _user: User = Depends(current_user)):
    rows = report_data(db, start, end)
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=["vehicle", "plate", "fuel_amount", "liters", "distance_km", "maintenance_cost", "total_cost", "km_per_liter", "cost_per_km"], delimiter=";")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row[key] for key in writer.fieldnames})
    return StreamingResponse(iter(["\ufeff" + output.getvalue()]), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=cpdesk-frota.csv"})


@router.get("/reports/fleet.xlsx")
def fleet_report_xlsx(start: date | None = None, end: date | None = None,
                      db: Session = Depends(get_db), _user: User = Depends(current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    rows = report_data(db, start, end)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Frota"
    columns = [("vehicle", "Veículo"), ("plate", "Placa"), ("fuel_amount", "Combustível (R$)"),
               ("liters", "Litros"), ("distance_km", "Distância (km)"),
               ("maintenance_cost", "Manutenção (R$)"), ("total_cost", "Custo total (R$)"),
               ("km_per_liter", "Km/l"), ("cost_per_km", "R$/km")]
    sheet.append([label for _, label in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25354A")
    for row in rows:
        sheet.append([float(row[key]) if isinstance(row[key], Decimal) else row[key] for key, _ in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=cpdesk-frota.xlsx"})


def _pdf_response(title: str, headers: list[str], rows: list[list[object]], filename: str):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=12 * mm,
        leftMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    table_data = [[Paragraph(f"<b>{escape(str(value))}</b>", styles["BodyText"]) for value in headers]]
    for row in rows:
        table_data.append([Paragraph(escape("—" if value is None else str(value)), styles["BodyText"]) for value in row])
    table = Table(table_data, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#26364d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d9dee8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f5f8")]),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    document.build([Paragraph(escape(title), styles["Title"]), Spacer(1, 5 * mm), table])
    output.seek(0)
    return StreamingResponse(output, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.get("/reports/fleet.pdf")
def fleet_report_pdf(start: date | None = None, end: date | None = None,
                     db: Session = Depends(get_db), _user: User = Depends(current_user)):
    rows = report_data(db, start, end)
    if start and end and start > end:
        raise HTTPException(422, "Data inicial posterior à data final")
    data = [[row["vehicle"], row["plate"], money_value(row["fuel_amount"]), number_value(row["liters"]),
        number_value(row["distance_km"]), money_value(row["maintenance_cost"]), money_value(row["total_cost"]),
        number_value(row["km_per_liter"]), money_value(row["cost_per_km"])] for row in rows]
    return _pdf_response("Relatório de despesas e consumo da frota",
        ["Veículo", "Placa", "Combustível", "Litros", "Distância", "Manutenção", "Total", "Km/l", "R$/km"],
        data, "cpdesk-frota.pdf")


def money_value(value):
    return "—" if value is None else f"R$ {Decimal(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def number_value(value):
    return "—" if value is None else f"{Decimal(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


@router.get("/reports/stock-consumption")
def stock_consumption_report(start: date | None = None, end: date | None = None,
                             db: Session = Depends(get_db), _user: User = Depends(current_user)):
    if start and end and start > end:
        raise HTTPException(422, "Data inicial posterior à data final")
    return stock_consumption(db, start, end)


@router.get("/reports/stock-consumption.xlsx")
def stock_consumption_xlsx(start: date | None = None, end: date | None = None,
                           db: Session = Depends(get_db), _user: User = Depends(current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    if start and end and start > end:
        raise HTTPException(422, "Data inicial posterior à data final")
    consumptions = stock_consumption(db, start, end)
    query = select(StockMovement, Item).join(Item).where(StockMovement.movement_type == "entry").order_by(StockMovement.created_at.desc())
    if start:
        query = query.where(StockMovement.created_at >= datetime.combine(start, datetime.min.time()))
    if end:
        query = query.where(StockMovement.created_at < datetime.combine(end + timedelta(days=1), datetime.min.time()))
    entries = [{"date": movement.created_at.date(), "code": item.sku, "item": item.name,
        "quantity": movement.quantity, "unit": item.unit, "unit_cost": movement.unit_cost,
        "invoice_total": movement.invoice_total, "supplier": movement.supplier_name,
        "invoice_number": movement.invoice_number, "reason": movement.reason}
        for movement, item in db.execute(query)]
    workbook = Workbook()
    sheets = [
        (workbook.active, "Consumo", consumptions, [("date", "Data"), ("code", "Código"), ("item", "Produto"),
            ("sector", "Setor"), ("responsible", "Funcionário"), ("quantity", "Quantidade"), ("unit", "Unidade"),
            ("unit_cost", "Custo unitário"), ("total_cost", "Custo estimado"), ("reason", "Motivo")]),
        (workbook.create_sheet("Entradas"), "Entradas", entries, [("date", "Data"), ("code", "Código"), ("item", "Produto"),
            ("quantity", "Quantidade"), ("unit", "Unidade"), ("unit_cost", "Valor unitário"),
            ("invoice_total", "Total da nota"), ("supplier", "Fornecedor"), ("invoice_number", "Nota fiscal"), ("reason", "Motivo")]),
    ]
    for sheet, name, rows, columns in sheets:
        sheet.title = name
        sheet.append([label for _, label in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="25354A")
        for row in rows:
            sheet.append([float(row[key]) if isinstance(row.get(key), Decimal) else row.get(key) for key, _ in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=cpdesk-consumo-produtos.xlsx"})


@router.get("/reports/stock-consumption.pdf")
def stock_consumption_pdf(start: date | None = None, end: date | None = None,
                          db: Session = Depends(get_db), _user: User = Depends(current_user)):
    if start and end and start > end:
        raise HTTPException(422, "Data inicial posterior à data final")
    rows = stock_consumption(db, start, end)
    data = [[row["date"].strftime("%d/%m/%Y"), row["code"] or "—", row["item"], row["sector"],
        row["responsible"], number_value(row["quantity"]), row["unit"], money_value(row["total_cost"])] for row in rows]
    return _pdf_response("Relatório de consumo de produtos por setor",
        ["Data", "Código", "Produto", "Setor", "Funcionário", "Quantidade", "Unid.", "Custo estimado"],
        data, "cpdesk-consumo-produtos.pdf")


@router.get("/reports/stock.csv")
def stock_report_csv(db: Session = Depends(get_db), _user: User = Depends(current_user)):
    balances = item_balances(db)
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Código", "Item", "Unidade", "Saldo", "Mínimo"])
    for item in db.scalars(select(Item).order_by(Item.name)).all():
        writer.writerow([item.sku or "", item.name, item.unit, balances.get(item.id, 0), item.minimum_stock])
    return StreamingResponse(iter(["\ufeff" + output.getvalue()]), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=cpdesk-estoque.csv"})
